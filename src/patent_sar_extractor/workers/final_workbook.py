"""Workbook/SDF presentation helpers; no identity selection or acceptance."""

from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from patent_sar_extractor.core.activity_identity import normalize_compound
from patent_sar_extractor.core.activity_join import activity_for_compound

from .final_activity_sheet import write_activity_observations
from .final_data import get_smiles_for_binding

GRADE_FILLS = {
    "A": PatternFill("solid", fgColor="92D050"),
    "B": PatternFill("solid", fgColor="FFC000"),
    "C": PatternFill("solid", fgColor="FF9900"),
    "D": PatternFill("solid", fgColor="FF4444"),
    "+": PatternFill("solid", fgColor="92D050"),
    "-": PatternFill("solid", fgColor="FF4444"),
}

# Excel样式
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
NOTE_FONT = Font(size=9, color="666666")
NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")

# 图片嵌入尺寸
IMG_WIDTH_PX = 450
IMG_HEIGHT_PX = 300
ROW_HEIGHT_PT = 230

# 固定列定义 (name, width) — 活性列在运行时动态生成
FIXED_COLUMNS = [
    ("Cpd ID", 14),
    ("正文化合物编号", 16),
    ("页码", 8),
    ("结构索引", 10),
    ("结构图", 62),
    ("SMILES", 55),
    ("分子式", 16),
    ("分子量", 10),
    ("环数", 6),
    ("手性中心", 8),
]


def resolve_display_image_path(binding, project_root):
    """Choose the human-facing structure image, not an OCSR repair crop."""
    for key in ("display_image_path", "source_image_path", "image_path"):
        path = str(binding.get(key, "") or "").strip()
        if not path:
            continue
        if not os.path.isabs(path):
            path = os.path.join(project_root, path)
        if os.path.exists(path):
            return path
    return ""


def build_assay_columns(assays):
    """从assay元数据动态生成活性列定义

    Returns:
        list of (column_header, width, target, sdf_prop_name)
    """
    columns = []
    for assay in assays:
        target = assay.get("target", "Unknown")
        test_type = assay.get("test_type", "")
        unit = assay.get("unit", "")

        # 生成列标题: "IGF-1R IC50" 或 "IR IC50 (nM)"
        header = target
        if test_type:
            header = f"{target} {test_type}"
        if unit:
            header = f"{header} ({unit})"

        # SDF属性名: Activity_IGF1R_IC50
        sdf_prop = (
            f"Activity_{target.replace('-', '').replace(' ', '_')}_{test_type}"
            if test_type
            else f"Activity_{target.replace('-', '')}"
        )

        columns.append((header, 14, target, sdf_prop))

    return columns


def build_grade_note(assays):
    """从assay元数据自动生成分级说明"""
    if not assays:
        return ""

    lines = []
    for assay in assays:
        target = assay.get("target", "")
        test_type = assay.get("test_type", "")
        grades = assay.get("grades", {})

        if not grades:
            continue

        header = f"{target} {test_type}" if test_type else target
        lines.append(f"【{header}】")
        for grade_key, grade_desc in sorted(grades.items()):
            lines.append(f"  {grade_key}: {grade_desc}")

    return "\n".join(lines) if lines else ""


def generate_excel(
    bindings,
    smiles_map,
    smiles_by_binding_key,
    smiles_by_structure_id,
    act_data,
    assays,
    output_path,
    project_root,
    activity_rows=None,
):
    """生成最终 Excel。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Final Results"

    # 构建列定义: 固定列 + 动态活性列 + 分级说明列
    assay_cols = build_assay_columns(assays)
    grade_note = build_grade_note(assays)

    all_columns = list(FIXED_COLUMNS) + [(h, w) for h, w, _, _ in assay_cols]
    if grade_note:
        all_columns.append(("活性分级说明", 40))

    # 表头
    for col_idx, (h, w) in enumerate(all_columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # 活性列起始位置
    activity_col_start = len(FIXED_COLUMNS) + 1

    # 数据行
    img_count = 0
    for row_idx, b in enumerate(bindings, 2):
        cpd = normalize_compound(b.get("cpd", b.get("cpd_id", "")))
        compound_id = normalize_compound(b.get("compound_id", "")) or cpd
        page = b.get("page_no", "")
        sidx = b.get("structure_index", b.get("struct_idx", ""))

        # 固定列
        ws.cell(row=row_idx, column=1, value=cpd).border = THIN_BORDER
        ws.cell(row=row_idx, column=2, value=compound_id).border = THIN_BORDER
        ws.cell(row=row_idx, column=3, value=page).border = THIN_BORDER
        ws.cell(row=row_idx, column=4, value=sidx).border = THIN_BORDER

        # 结构图
        img_path = resolve_display_image_path(b, project_root)
        if img_path and os.path.exists(img_path):
            try:
                img = XlImage(img_path)
                img.width = IMG_WIDTH_PX
                img.height = IMG_HEIGHT_PX
                ws.add_image(img, f"E{row_idx}")
                img_count += 1
            except (OSError, ValueError, TypeError) as e:
                ws.cell(
                    row=row_idx, column=5, value=f"图片加载失败: {e}"
                ).border = THIN_BORDER
        else:
            ws.cell(
                row=row_idx,
                column=5,
                value="" if b.get("activity_only") else "图片未找到",
            ).border = THIN_BORDER

        # SMILES
        sr = get_smiles_for_binding(
            b, smiles_map, smiles_by_binding_key, smiles_by_structure_id
        )
        ws.cell(
            row=row_idx, column=6, value=sr.get("canonical_smiles", "")
        ).border = THIN_BORDER
        ws.cell(row=row_idx, column=6).alignment = Alignment(wrap_text=True)
        ws.cell(
            row=row_idx, column=7, value=str(sr.get("mol_formula", ""))
        ).border = THIN_BORDER
        ws.cell(
            row=row_idx, column=8, value=str(sr.get("mol_weight", ""))
        ).border = THIN_BORDER
        ws.cell(
            row=row_idx, column=9, value=sr.get("ring_count", "")
        ).border = THIN_BORDER
        ws.cell(
            row=row_idx, column=10, value=sr.get("chiral_centers", "")
        ).border = THIN_BORDER

        # 动态活性列
        act = activity_for_compound(act_data, cpd)
        for i, (header, width, target, sdf_prop) in enumerate(assay_cols):
            col = activity_col_start + i
            grade = act.get(target, "") if isinstance(act, dict) else ""
            cell = ws.cell(row=row_idx, column=col, value=grade)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER
            cell.font = Font(bold=True, size=14)
            if grade in GRADE_FILLS:
                cell.fill = GRADE_FILLS[grade]

        # 活性分级说明
        if grade_note:
            note_col = activity_col_start + len(assay_cols)
            note_cell = ws.cell(row=row_idx, column=note_col, value=grade_note)
            note_cell.font = NOTE_FONT
            note_cell.fill = NOTE_FILL
            note_cell.alignment = Alignment(wrap_text=True, vertical="center")
            note_cell.border = THIN_BORDER

        ws.row_dimensions[row_idx].height = ROW_HEIGHT_PT

    if act_data:
        act_ws = wb.create_sheet("Activity Results")
        activity_targets = []
        for values in act_data.values():
            if not isinstance(values, dict):
                continue
            for target in values:
                if target not in activity_targets:
                    activity_targets.append(target)
        act_headers = ["Cpd ID"] + activity_targets
        for col_idx, header in enumerate(act_headers, 1):
            cell = act_ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = THIN_BORDER
            act_ws.column_dimensions[get_column_letter(col_idx)].width = (
                18 if col_idx == 1 else 24
            )

        for row_idx, cpd in enumerate(act_data, 2):
            values = (
                act_data.get(cpd, {}) if isinstance(act_data.get(cpd), dict) else {}
            )
            act_ws.cell(row=row_idx, column=1, value=cpd).border = THIN_BORDER
            for col_idx, target in enumerate(activity_targets, 2):
                act_ws.cell(
                    row=row_idx, column=col_idx, value=values.get(target, "")
                ).border = THIN_BORDER

    write_activity_observations(wb, activity_rows or [])
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    wb.save(output_path)
    return img_count

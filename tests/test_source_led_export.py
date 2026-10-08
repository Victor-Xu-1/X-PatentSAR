"""Real Excel/SDF and deterministic QA on named synthetic controls, not model accuracy."""

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from PIL import Image

from patent_sar_extractor import contracts
from patent_sar_extractor.core.activity_coverage import coverage_packet, source_record
from patent_sar_extractor.core.binding_artifacts import write_binding_result
from patent_sar_extractor.core.formal_structure import (
    FORMAL_SCOPE,
    SOURCE_EXECUTION_MODE,
)
from patent_sar_extractor.core.ocsr.smiles_qc import qc_smiles
from patent_sar_extractor.core.ocsr.stereo_evidence import check_source_stereochemistry
from patent_sar_extractor.core.qa_report import build_qa_report
from patent_sar_extractor.smiles_artifact import build_smiles_artifact
from tests.test_source_led_pipeline import proved_binding


def save(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data), encoding="utf-8")


def control_record(binding, raw="CCO"):
    # An explicit synthetic graph oracle, never a pretend recognizer inference.
    digest = hashlib.sha256(Path(binding["image_path"]).read_bytes()).hexdigest()
    qc = qc_smiles(raw)
    evidence = check_source_stereochemistry(
        qc,
        {
            "version": contracts.STEREO_EVIDENCE_VERSION,
            "image_sha256": digest,
            "image_size": [20, 20],
            "unknown_bond_boxes": [],
        },
    )
    return {
        "cpd_id": binding["cpd"],
        "structure_id": binding["structure_id"],
        "image_hash": digest,
        "raw_smiles": raw,
        "canonical_smiles": qc["canonical_smiles"],
        "rdkit_valid": qc["rdkit_valid"],
        "OCSR_quality_flag": qc["quality_flag"],
        "stereochemistry": evidence,
    }


def run_fixture(root, *, activities=None, invalid_second=False, classified_pages=None):
    # Explicit synthetic source inventory, not proof of a real patent/model run.
    seeds = (
        classified_pages
        if classified_pages is not None
        else ([] if not activities else [0])
    )
    regions = (
        [
            source_record(
                1,
                "text",
                None,
                "Controlled table",
                "Declared control values",
                "parsed",
                "declared_text",
                len(activities),
            )
        ]
        if activities
        else []
    )
    coverage = coverage_packet(
        seeds, regions, [SimpleNamespace(page_no=1) for _ in activities or []]
    )
    for number in (1, 2):
        Image.new("RGB", (20, 20), "white").save(root / f"control-{number}.png")
    bindings = [
        proved_binding("Compound 42", "S42", str(root / "control-1.png")),
        proved_binding("Compound 42A", "S42A", str(root / "control-2.png")),
    ]
    payload = write_binding_result(
        root / "structure_bindings",
        patent_id="CONTROL",
        bindings=bindings,
        detected_style="synthetic_control",
        include_intermediates=False,
        total_structures=2,
        total_compound_blocks=0,
        table_pages=[],
        table_covered_count=0,
        no_binding=[],
        unbound_pages=[],
    )
    bindings = payload["final_bindings"]
    records = [control_record(binding) for binding in bindings]
    if invalid_second:
        records[1]["OCSR_quality_flag"] = "stereo_source_conflict"
    save(
        root / "smiles/smiles_results.json",
        {
            **build_smiles_artifact(records),
            "formal_acceptance_scope": FORMAL_SCOPE,
            "binding_execution_mode": SOURCE_EXECUTION_MODE,
        },
    )
    save(
        root / "activity/activity_data.json",
        {
            **contracts.artifact_identity(
                contracts.ACTIVITY_SCHEMA, contracts.ACTIVITY_SCHEMA_VERSION
            ),
            "rows": activities or [],
            "coverage": coverage,
        },
    )
    save(
        root / "page_classification/page_classification.json",
        {
            **contracts.artifact_identity(
                contracts.PAGE_CLASSIFICATION_SCHEMA,
                contracts.PAGE_CLASSIFICATION_SCHEMA_VERSION,
            ),
            "page_count": 1,
            "activity_pages": classified_pages
            if classified_pages is not None
            else ([] if not activities else [0]),
        },
    )
    save(root / "structures/metadata.json", {"total_structures": 2})
    save(root / "structure_pages/locator.json", {})
    save(
        root / "pipeline_summary.json",
        {
            **contracts.artifact_identity(
                contracts.RUN_SUMMARY_SCHEMA, contracts.RUN_SUMMARY_SCHEMA_VERSION
            ),
            "status": "qa_pending",
            "main_chain": list(contracts.CORE_STAGE_ORDER),
            "patent_id": "CONTROL",
            "steps": {},
        },
    )
    return bindings


def export(root, *, continue_scientific=False):
    worker = (
        Path(__file__).parents[1]
        / "src/patent_sar_extractor/workers/gen_final_results.py"
    )
    command = [
        sys.executable,
        "-B",
        str(worker),
        "--bindings",
        str(root / "structure_bindings/bindings.json"),
        "--smiles",
        str(root / "smiles/smiles_results.json"),
        "--activity",
        str(root / "activity/activity_data.json"),
        "--classification",
        str(root / "page_classification/page_classification.json"),
        "--output-dir",
        str(root / "final_results"),
        "--patent",
        "CONTROL",
        "--rdkit-python",
        sys.executable,
    ]
    if continue_scientific:
        command.append("--continue-on-scientific-errors")
    return subprocess.run(
        command, capture_output=True, text=True, timeout=30, check=False
    )


class SourceLedExportTests(unittest.TestCase):
    def test_review_workbook_keeps_all_ids_images_and_activity_without_bad_chemistry(
        self,
    ):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            activities = [
                {"cpd": "Compound 42", "activity_values": {"IC50 (nM)": "17"}},
                {"cpd": "Compound 42A", "activity_values": {"IC50 (nM)": "29"}},
            ]
            run_fixture(root, activities=activities, invalid_second=True)
            result = export(root, continue_scientific=True)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            workbook = load_workbook(root / "final_results/CONTROL_final.xlsx")
            sheet = workbook["Final Results"]
            rows = list(sheet.values)
            self.assertEqual(
                [row[0] for row in rows[1:]], ["Compound 42", "Compound 42A"]
            )
            activity_column = rows[0].index("IC50 (nM)")
            status_column = rows[0].index("识别状态")
            self.assertEqual([row[activity_column] for row in rows[1:]], ["17", "29"])
            self.assertEqual(rows[1][status_column], "当前规则通过")
            self.assertEqual(rows[2][status_column], "待复核")
            self.assertTrue(all(value is None for value in rows[2][5:10]))
            self.assertEqual(len(sheet._images), 2)
            self.assertTrue(
                all(
                    image.anchor.ext.cx == image.anchor.ext.cy
                    for image in sheet._images
                )
            )
            self.assertEqual(sheet.freeze_panes, "F2")
            self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
            receipt = json.loads(
                (root / "final_results/export_validation.json").read_text()
            )
            self.assertFalse(receipt["formal_acceptance_authority"])
            self.assertEqual(receipt["workbook_review_cpds"], ["Compound 42A"])
            self.assertEqual(
                receipt["strict_coverage"]["workbook_cpds"],
                ["Compound 42", "Compound 42A"],
            )
            self.assertEqual(
                receipt["strict_coverage"]["exported_cpds"], ["Compound 42"]
            )
            self.assertEqual(
                (root / "final_results/CONTROL_final.sdf").read_text().count("$$$$"), 1
            )
            qa = build_qa_report(str(root), patent_id="CONTROL")
            self.assertFalse(qa["acceptance"]["ok"])
            self.assertFalse(
                any(
                    "Excel main sheet does not match" in item
                    or "missing or changed activity" in item
                    for item in qa["acceptance"]["hard_errors"]
                )
            )

    def test_proved_no_activity_exports_all_ids_and_passes_control_qa(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            run_fixture(root)
            result = export(root)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            workbook = load_workbook(root / "final_results/CONTROL_final.xlsx")
            self.assertEqual(
                [row[0] for row in workbook["Final Results"].values][1:],
                ["Compound 42", "Compound 42A"],
            )
            self.assertNotIn("Activity Results", workbook.sheetnames)
            qa = build_qa_report(str(root), patent_id="CONTROL")
            self.assertTrue(qa["acceptance"]["ok"], qa["acceptance"]["hard_errors"])
            self.assertEqual(qa["acceptance"]["sdf_record_count"], 2)
            self.assertEqual(qa["formal_acceptance_scope"], FORMAL_SCOPE)

    def test_activity_is_left_join_and_repeated_values_are_not_overwritten(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            activities = [
                {"cpd": "Compound 42", "activity_values": {"IC50 (nM)": "17"}},
                {"cpd": "Compound 42", "activity_values": {"IC50 (nM)": "29"}},
            ]
            run_fixture(root, activities=activities)
            result = export(root)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            workbook = load_workbook(root / "final_results/CONTROL_final.xlsx")
            rows = list(workbook["Final Results"].values)
            activity_column = rows[0].index("IC50 (nM)")
            self.assertEqual(rows[1][activity_column], "17 | 29")
            self.assertIsNone(rows[2][activity_column])
            source_rows = list(workbook["Activity Observations"].values)
            self.assertEqual([row[2] for row in source_rows[1:]], ["17", "29"])
            self.assertTrue(all(row[3] is None for row in source_rows[1:]))
            qa = build_qa_report(str(root), patent_id="CONTROL")
            self.assertTrue(qa["acceptance"]["ok"], qa["acceptance"]["hard_errors"])

    def test_bad_science_keeps_failure_and_exports_only_qualified_exact_pairs(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            run_fixture(root, invalid_second=True)
            result = export(root, continue_scientific=True)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            receipt = json.loads(
                (root / "final_results/export_validation.json").read_text()
            )
            self.assertTrue(receipt["hard_errors"])
            self.assertEqual(
                receipt["strict_coverage"]["expected_cpds"],
                ["Compound 42", "Compound 42A"],
            )
            self.assertEqual(
                receipt["strict_coverage"]["exported_cpds"], ["Compound 42"]
            )
            self.assertTrue((root / "STRICT_ACCEPTANCE_FAILED.json").exists())
            qa = build_qa_report(
                str(root), patent_id="CONTROL", ignore_previous_failure_marker=True
            )
            self.assertFalse(qa["acceptance"]["ok"])
            self.assertTrue(qa["acceptance"]["hard_errors"])

    def test_zero_qualified_structures_still_publish_rejected_research_packet(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            run_fixture(root)
            smiles = root / "smiles/smiles_results.json"
            packet = json.loads(smiles.read_text())
            for record in packet["records"]:
                record["OCSR_quality_flag"] = "stereo_source_conflict"
            save(smiles, packet)
            result = export(root, continue_scientific=True)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            receipt = json.loads(
                (root / "final_results/export_validation.json").read_text()
            )
            self.assertEqual(receipt["strict_coverage"]["exported_cpds"], [])
            self.assertTrue(receipt["hard_errors"])
            self.assertTrue((root / "final_results/CONTROL_final.xlsx").is_file())
            self.assertEqual(
                (root / "final_results/CONTROL_final.sdf").stat().st_size, 0
            )
            qa = build_qa_report(
                str(root), patent_id="CONTROL", ignore_previous_failure_marker=True
            )
            self.assertFalse(qa["acceptance"]["ok"])

    def test_declared_activity_pages_lost_rows_cannot_be_accepted(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            run_fixture(root, classified_pages=[0])
            result = export(root)
            self.assertNotEqual(result.returncode, 0)
            result = export(root, continue_scientific=True)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            qa = build_qa_report(
                str(root), patent_id="CONTROL", ignore_previous_failure_marker=True
            )
            self.assertFalse(qa["acceptance"]["ok"])
            self.assertIn(
                "classification proof", " ".join(qa["acceptance"]["hard_errors"])
            )

    def test_malformed_ownership_hard_stops_even_with_continue(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            run_fixture(root)
            smiles = root / "smiles/smiles_results.json"
            payload = json.loads(smiles.read_text())
            payload["records"][1]["structure_id"] = "S42"
            save(smiles, payload)
            result = export(root, continue_scientific=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse((root / "final_results/CONTROL_final.xlsx").exists())

    def test_persisted_earlier_scientific_failure_cannot_be_cleared_by_good_exports(
        self,
    ):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            run_fixture(root)
            self.assertEqual(export(root).returncode, 0)
            summary = json.loads((root / "pipeline_summary.json").read_text())
            summary["scientific_errors"] = {
                "bind": ["known withheld printed identifier"]
            }
            save(root / "pipeline_summary.json", summary)
            qa = build_qa_report(
                str(root), patent_id="CONTROL", ignore_previous_failure_marker=True
            )
            self.assertFalse(qa["acceptance"]["ok"])
            self.assertIn(
                "bind: known withheld printed identifier",
                qa["acceptance"]["hard_errors"],
            )


if __name__ == "__main__":
    unittest.main()
